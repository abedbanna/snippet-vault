#!/usr/bin/env python3
"""Write the Day 2 Session 4 scoring workbooks: an empty M4-scoring template and the filled leave-process example.

Columns follow Module 4's lab step 4. Value, risk total and quadrant are live formulas, so the files work the same after
import into Google Sheets. The example rows are make_day2s4_figures.CANDIDATES (the matrix slide plots the same rows).
Usage: python3 tools/make_day2s4_sheets.py out/day2/Session4/03_Participant_Handouts
"""
import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from make_day2s4_figures import CANDIDATES, VALUE_HIGH, RISK_HIGH  # noqa: E402

HEAD = ["Candidate", "Steps included", "Minutes saved per run", "Runs per month", "Value (hours/month)",
        "Reversibility 0-3", "Money 0-3", "People 0-3", "Legal/data 0-3", "Risk total", "Quadrant", "Mode", "Decision", "Notes"]
WIDTH = [44, 16, 12, 11, 12, 12, 10, 10, 11, 10, 18, 26, 44, 40]
BOLD, THIN = Font(bold=True), Side(style="thin", color="D1D5DB")
ROWS = 12


def sheet(ws, rows):
    ws.append(HEAD)
    for c, w in zip(ws[1], WIDTH):
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="111111")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 34
    for i in range(ROWS):
        r = i + 2
        data = rows[i] if i < len(rows) else None
        vals = [data[0], data[1], data[2], data[3]] if data else [None] * 4
        risk = list(data[4:8]) if data else [None] * 4
        ws.cell(r, 1, vals[0]); ws.cell(r, 2, vals[1]); ws.cell(r, 3, vals[2]); ws.cell(r, 4, vals[3])
        ws.cell(r, 5, f'=IF(OR(C{r}="",D{r}=""),"",ROUND(C{r}*D{r}/60,1))')
        for j, v in enumerate(risk):
            ws.cell(r, 6 + j, v)
        ws.cell(r, 10, f'=IF(COUNT(F{r}:I{r})<4,"",SUM(F{r}:I{r}))')
        ws.cell(r, 11, f'=IF(OR(E{r}="",J{r}=""),"",IF(E{r}>={VALUE_HIGH:g},IF(J{r}<{RISK_HIGH},"Build first","Build with a gate"),'
                       f'IF(J{r}<{RISK_HIGH},"Batch later","Leave alone")))')
        ws.cell(r, 12, f'=IF(J{r}="","",IF(J{r}<=3,"Fully automatic",IF(J{r}<=7,"Notify a person, with undo","Approval gate every time")))')
        ws.cell(r, 13, data[8] if data else None)
        for c in ws[r]:
            c.border = Border(bottom=THIN); c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in (5, 10, 11, 12):
            ws.cell(r, col).font = BOLD
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="3", showErrorMessage=True,
                        error="Score each risk question 0, 1, 2 or 3")
    ws.add_data_validation(dv); dv.add(f"F2:I{ROWS + 1}")
    fills = {"Build first": "DCFCE7", "Build with a gate": "FEF3C7", "Batch later": "F3F4F6", "Leave alone": "FEE2E2"}
    for txt, col in fills.items():
        ws.conditional_formatting.add(f"K2:K{ROWS + 1}", FormulaRule(formula=[f'K2="{txt}"'], fill=PatternFill("solid", fgColor=col)))
    ws.freeze_panes = "B2"
    r = ROWS + 3
    notes = [f"Value = minutes saved per run x runs per month / 60. Count only the minutes the automation removes; subtract any review time.",
             f"Risk = the four scores added (0 to 12). Where two people disagree, take the higher score and write why in Notes.",
             f"Quadrant: high value is {VALUE_HIGH:g} hours a month or more; high risk is {RISK_HIGH} or more.",
             "Mode: 0 to 3 fully automatic; 4 to 7 runs, a person is notified and can undo; 8 to 12 stops at an approval gate."]
    for t in notes:
        ws.cell(r, 1, t).font = Font(italic=True, color="6B7280"); r += 1
    return r + 1


def roi(ws, r):
    ws.cell(r, 1, "ROI range for the first build (candidate 1)").font = Font(bold=True, size=12); r += 1
    lines = [("", "Low", "High"), ("Requests per month", 50, 70), ("Minutes saved per request", 5, 7),
             ("Hours saved per month", None, None), ("Loaded hourly cost (JOD)", 10, 10),
             ("Gross value per month (JOD)", None, None), ("Tool cost per month (JOD)", 15, 15),
             ("Maintenance: 1 hour a month (JOD)", 10, 10), ("Net value per month (JOD)", None, None)]
    top = r
    for i, (a, b, c) in enumerate(lines):
        ws.cell(r, 1, a); ws.cell(r, 2, b); ws.cell(r, 3, c)
        if i == 0:
            ws.cell(r, 2).font = ws.cell(r, 3).font = BOLD
        r += 1
    for col in "BC":
        ws[f"{col}{top + 3}"] = f"=ROUND({col}{top + 1}*{col}{top + 2}/60,1)"
        ws[f"{col}{top + 5}"] = f"={col}{top + 3}*{col}{top + 4}"
        ws[f"{col}{top + 8}"] = f"={col}{top + 5}-{col}{top + 6}-{col}{top + 7}"
        for k in (3, 5, 8):
            ws[f"{col}{top + k}"].font = BOLD
    ws.cell(r, 1, "Say it as: about 4 to 8 hours a month, worth JOD 15 to 60 net, assuming 50 to 70 requests and 5 to 7 "
                  "minutes saved each.").font = Font(italic=True)
    return r + 1


def main(out):
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "M4-scoring"
    sheet(ws, []); wb.save(out / "Day2_S4_M4_Scoring_Template.xlsx")
    wb = Workbook(); ws = wb.active; ws.title = "M4-scoring"
    r = sheet(ws, CANDIDATES); roi(ws, r)
    ws.cell(2, 14, "Chosen first build. Gate: the manager's approval, before the leave is recorded.")
    ws.cell(5, 14, "Low value by this measure, but it is what fixes the waiting (the 4 days)")
    wb.save(out / "Day2_S4_M4_Scoring_Example.xlsx")
    print("wrote", out / "Day2_S4_M4_Scoring_Template.xlsx", out / "Day2_S4_M4_Scoring_Example.xlsx")


if __name__ == "__main__":
    main(pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "out/day2/Session4/03_Participant_Handouts"))
