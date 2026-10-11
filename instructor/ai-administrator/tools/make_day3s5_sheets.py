"""Day 3 Session 5 spreadsheets: the leave-demo sheet (Balances + Requests) and the Module 05 lab template (Requests).
Upload to Google Drive and open with Google Sheets; the tab names and headers must stay exactly as written."""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
out = sys.argv[1] if len(sys.argv) > 1 else '.'
os.makedirs(out, exist_ok=True)
HEAD = Font(bold=True, color='FFFFFF'); FILL = PatternFill('solid', fgColor='0F766E')
def header(ws, cols, widths=None):
    ws.append(cols)
    for i, c in enumerate(ws[1], 1):
        c.font = HEAD; c.fill = FILL; c.alignment = Alignment(vertical='center')
        ws.column_dimensions[c.column_letter].width = (widths or {}).get(cols[i-1], max(14, len(cols[i-1]) + 4))
    ws.freeze_panes = 'A2'

# 1. Leave demo. Replace the first email with your own so the demo finds your balance.
wb = Workbook(); b = wb.active; b.title = 'Balances'
header(b, ['Employee', 'Email', 'Department', 'Annual entitlement', 'Taken this year', 'Remaining days'], {'Email': 30})
for row in [('Demo Instructor', 'YOUR-EMAIL@gmail.com', 'Operations', 24, 12),
            ('Lina Haddad', 'lina.haddad@aicaravan.example', 'Finance', 24, 22),
            ('Omar Khalil', 'omar.khalil@aicaravan.example', 'IT', 24, 3),
            ('Sara Nasser', 'sara.nasser@aicaravan.example', 'HR', 24, 9)]:
    b.append(list(row)); r = b.max_row
    b.cell(r, 6).value = f'=D{r}-E{r}'
r = wb.create_sheet('Requests')
header(r, ['Request ID', 'Submitted', 'Employee', 'Email', 'Manager email', 'Leave type', 'First day', 'Last day', 'Working days',
           'Notice (working days)', 'Notice needed', 'Remaining before', 'Quarter end', 'Reason', 'Status'], {'Reason': 44, 'Email': 30, 'Manager email': 30})
wb.save(os.path.join(out, 'Day3_S5_Leave_Demo_Sheet.xlsx'))

# 2. Module 05 lab template
wb = Workbook(); r = wb.active; r.title = 'Requests'
header(r, ['Timestamp', 'Name', 'Email', 'Department', 'Amount', 'Description', 'Status'], {'Email': 30, 'Description': 40})
wb.save(os.path.join(out, 'Day3_S5_Requests_Template.xlsx'))
print('wrote 2 workbooks to', out)
