#!/usr/bin/env python3
"""Build a Kahoot import spreadsheet (+ a markdown answer key) from a questions YAML, and enforce anti-guessing rules.

Usage: build_kahoot.py content/day1_kahoot_session1.yaml out/day1/Day1_S1_Kahoot.xlsx [out/day1/Day1_S1_Kahoot_Answer_Key.md]

Rules enforced (fail the build if broken):
  - question <= 120 chars, each answer <= 75 chars (Kahoot limits)
  - exactly 4 options, one correct
  - option lengths balanced: the correct option is never the longest by more than SLACK chars, and the spread
    between the longest and shortest option is <= SPREAD chars
  - correct positions distributed: no position holds more than ceil(n/4)+1 answers
  - no "all of the above" / "none of the above"
Spreadsheet layout follows Kahoot's template: header row 8, data from row 9, columns B..H.
"""
import sys, math, pathlib, yaml
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

SLACK = 6; SPREAD = 18
TIMES = {5, 10, 20, 30, 60, 90, 120, 240}


def validate(qs):
    errs = []; pos = [0, 0, 0, 0]
    for i, q in enumerate(qs, 1):
        o = q["options"]; a = q["answer"]
        if len(q["q"]) > 120: errs.append(f"Q{i}: question {len(q['q'])} chars > 120")
        if len(o) != 4: errs.append(f"Q{i}: needs 4 options")
        for j, x in enumerate(o):
            if len(x) > 75: errs.append(f"Q{i}: option {j+1} {len(x)} chars > 75")
            if "of the above" in x.lower(): errs.append(f"Q{i}: option {j+1} is an 'of the above'")
        lens = [len(x) for x in o]
        if lens[a] > max(l for k, l in enumerate(lens) if k != a) + SLACK:
            errs.append(f"Q{i}: correct option is the longest by > {SLACK} chars ({lens})")
        if max(lens) - min(lens) > SPREAD:
            errs.append(f"Q{i}: option length spread {max(lens)-min(lens)} > {SPREAD} ({lens})")
        pos[a] += 1
    cap = math.ceil(len(qs) / 4) + 1
    if max(pos) > cap: errs.append(f"correct-answer positions unbalanced: {pos} (cap {cap})")
    return errs, pos


def build(yaml_path, xlsx_path, md_path=None):
    doc = yaml.safe_load(pathlib.Path(yaml_path).read_text()); qs = doc["questions"]; meta = doc["meta"]
    errs, pos = validate(qs)
    if errs: print("\n".join(errs)); sys.exit(1)
    t = int(meta.get("time", 20)); assert t in TIMES
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
    ws["B1"] = "Kahoot! quiz import template"; ws["B1"].font = Font(bold=True, size=14)
    ws["B3"] = meta["title"]; ws["B4"] = meta["subtitle"]
    ws["B6"] = "Do not edit row 8. Questions start on row 9. Answers 1-4 in columns C-F; time limit in G; correct answer number(s) in H."
    hdr = ["Question - max 120 characters", "Answer 1 - max 75 characters", "Answer 2 - max 75 characters",
           "Answer 3 - max 75 characters", "Answer 4 - max 75 characters",
           "Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs", "Correct answer(s) - choose at least one"]
    for c, h in enumerate(hdr, start=2):
        cell = ws.cell(row=8, column=c, value=h); cell.font = Font(bold=True); cell.alignment = Alignment(wrap_text=True)
    for i, q in enumerate(qs):
        r = 9 + i
        ws.cell(row=r, column=2, value=q["q"])
        for j, o in enumerate(q["options"]): ws.cell(row=r, column=3 + j, value=o)
        ws.cell(row=r, column=7, value=int(q.get("time", t)))
        ws.cell(row=r, column=8, value=q["answer"] + 1)
    for col, w in zip("BCDEFGH", (70, 30, 30, 30, 30, 14, 14)): ws.column_dimensions[col].width = w
    pathlib.Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True); wb.save(xlsx_path)
    print(f"{xlsx_path}: {len(qs)} questions; correct positions A-D = {pos}")
    if md_path:
        L = "ABCD"; lines = ["---", f'title: "{meta["title"]}"', f'subtitle: "{meta["subtitle"]}"',
                             'kind: "Instructor answer key for the Kahoot import file"', "---", "",
                             "## How to import", "",
                             "1. In Kahoot, choose **Create**, then **Import spreadsheet** (the 'Add question' menu offers 'Import spreadsheet').",
                             f"2. Upload **{pathlib.Path(xlsx_path).name}**. All {len(qs)} questions arrive as quiz questions with a {t}-second timer.",
                             "3. Set the kahoot's cover title and run it in the room at the end of the session, or as a homework challenge.", "",
                             "## Design notes", "", meta.get("rules", "").strip(), "",
                             f"Correct answers by position A, B, C, D: {pos[0]}, {pos[1]}, {pos[2]}, {pos[3]}.", "",
                             "## Questions and answers", "", "| # | Question | A | B | C | D | Correct |", "|---|---|---|---|---|---|---|"]
        for i, q in enumerate(qs, 1):
            o = q["options"]; lines.append(f"| {i} | {q['q']} | {o[0]} | {o[1]} | {o[2]} | {o[3]} | **{L[q['answer']]}** |")
        pathlib.Path(md_path).write_text("\n".join(lines) + "\n"); print(md_path)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
