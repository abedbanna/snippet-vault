#!/usr/bin/env python3
"""Build a 16:9 PPTX deck from a content YAML.

Usage: build_deck.py content/day1_session1.yaml out/day1/Day1_S1_Foundations.pptx

Design rules (user request): white background, black text. Accent colour only on thin rules and small markers.
Figures come from figures/<slug>.png|jpg (see tools/fetch_figures.py); logos from brand/.
"""
import sys, json, pathlib, yaml
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = pathlib.Path(__file__).resolve().parent.parent
FIG = ROOT / "figures"; BRAND = ROOT / "brand"
MANIFEST = json.loads((FIG / "manifest.json").read_text())

BLACK = RGBColor(0x11, 0x11, 0x11); GREY = RGBColor(0x66, 0x66, 0x66); LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
RULE = RGBColor(0xDD, 0xDD, 0xDD); ACCENT = RGBColor(0xF7, 0xA1, 0x1A); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"; MONO = "Courier New"
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6); CONTENT_TOP = Inches(1.35); CONTENT_BOTTOM = H - Inches(0.75)
used_figures = set()


def rgb(run_or_font, color): run_or_font.color.rgb = color


def add_text(slide, left, top, width, height, text, size=18, bold=False, color=BLACK, align=PP_ALIGN.LEFT,
             font=FONT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(left, top, width, height); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        f = r.font; f.name = font; f.size = Pt(size); f.bold = bold; f.italic = italic; rgb(f, color)
    return tb


def add_bullets(slide, left, top, width, height, items, size=20, color=BLACK, space_after=8, numbered=False):
    tb = slide.shapes.add_textbox(left, top, width, height); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        marker = f"{i+1}.  " if numbered else "•  "
        r = p.add_run(); r.text = marker + item
        f = r.font; f.name = FONT; f.size = Pt(size); rgb(f, color)
    return tb


def fit_size(items, base, width_in, height_in, line_h=1.35):
    """Crude text fitting: shrink font until estimated wrapped lines fit the box."""
    size = base
    while size > 11:
        chars_per_line = max(10, int(width_in * 72 / (size * 0.5)))
        lines = sum(max(1, -(-len(t) // chars_per_line)) for t in items)
        if lines * size * line_h / 72 <= height_in: break
        size -= 1
    return size


def rule(slide, top, left=MARGIN, width=None, color=RULE, h=Pt(1.5)):
    width = width or (W - 2 * MARGIN)
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = color; shp.line.fill.background(); shp.shadow.inherit = False
    return shp


def footer(slide, meta, n):
    rule(slide, H - Inches(0.6))
    add_text(slide, MARGIN, H - Inches(0.55), Inches(10), Inches(0.4), meta["footer"], size=10, color=GREY)
    add_text(slide, W - MARGIN - Inches(1.5), H - Inches(0.55), Inches(1.5), Inches(0.4), str(n), size=10, color=GREY,
             align=PP_ALIGN.RIGHT)


def title_bar(slide, text, size=30):
    add_text(slide, MARGIN, Inches(0.4), W - 2 * MARGIN, Inches(0.85), text, size=size, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    rule(slide, Inches(1.22), width=Inches(1.2), color=ACCENT, h=Pt(3))
    rule(slide, Inches(1.22), left=MARGIN + Inches(1.2), width=W - 2 * MARGIN - Inches(1.2))


def notes(slide, text):
    if text: slide.notes_slide.notes_text_frame.text = text.strip()


def fig_path(slug):
    info = MANIFEST.get(slug)
    if not info: raise SystemExit(f"unknown figure slug: {slug}")
    used_figures.add(slug); return FIG / info["file"]


def place_image(slide, path, left, top, max_w, max_h):
    with Image.open(path) as im: w, h = im.size
    scale = min(max_w / w, max_h / h); pw, ph = int(w * scale), int(h * scale)
    return slide.shapes.add_picture(str(path), left + (max_w - pw) // 2, top + (max_h - ph) // 2, pw, ph)


def credit(slide, meta):
    add_text(slide, W - MARGIN - Inches(5), H - Inches(0.92), Inches(5), Inches(0.3), meta["figure_credit"], size=9,
             color=GREY, align=PP_ALIGN.RIGHT, italic=True)


def corner_logo(sl, s, top=Inches(0.3), h=Inches(0.85), w=Inches(2.6)):
    if s.get("logo") and (ROOT / s["logo"]).exists():
        place_image(sl, ROOT / s["logo"], W - MARGIN - w, top, w, h)


# ------------------------------------------------------------------ layouts
def l_title(prs, s, meta, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    place_image(sl, BRAND / "header_strip.png", MARGIN, Inches(0.45), W - 2 * MARGIN, Inches(1.1))
    rule(sl, Inches(3.0), width=Inches(1.6), color=ACCENT, h=Pt(4))
    add_text(sl, MARGIN, Inches(3.15), W - 2 * MARGIN, Inches(1.4), s["title"], size=40, bold=True)
    add_text(sl, MARGIN, Inches(4.55), W - 2 * MARGIN, Inches(0.9), s.get("subtitle", ""), size=20, color=GREY)
    add_text(sl, MARGIN, Inches(5.6), W - 2 * MARGIN, Inches(0.5), "Lead instructors: " + ", ".join(meta["instructors"]), size=16)
    add_text(sl, MARGIN, Inches(6.05), W - 2 * MARGIN, Inches(0.5),
             f"{meta['duration_min']} minutes · Module online: {meta['module_url'].replace('https://', '')}", size=13, color=GREY)
    footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_section(prs, s, meta, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(sl, MARGIN, Inches(2.2), W - 2 * MARGIN, Inches(0.5), f"Section · minute {s.get('minute', '')}" if 'minute' in s else "Section",
             size=14, color=GREY)
    rule(sl, Inches(2.75), width=Inches(1.6), color=ACCENT, h=Pt(4))
    add_text(sl, MARGIN, Inches(2.9), W - 2 * MARGIN, Inches(1.5), s["title"], size=38, bold=True)
    add_text(sl, MARGIN, Inches(4.4), W - 2 * MARGIN, Inches(0.9), s.get("subtitle", ""), size=20, color=GREY)
    if "duration" in s:
        add_text(sl, MARGIN, Inches(5.3), W - 2 * MARGIN, Inches(0.5), f"About {s['duration']} minutes", size=14, color=GREY)
    if s.get("logo") and (ROOT / s["logo"]).exists():   # big logo on section slides
        place_image(sl, ROOT / s["logo"], W - MARGIN - Inches(5.2), Inches(2.4), Inches(5.2), Inches(1.8))
    footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_bullets(prs, s, meta, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"]); corner_logo(sl, s)
    box_w, box_h = W - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP
    size = fit_size(s["bullets"], 22, box_w / 914400, box_h / 914400, line_h=1.6)
    add_bullets(sl, MARGIN, CONTENT_TOP + Inches(0.15), box_w, box_h, s["bullets"], size=size, space_after=12)
    footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_figure(prs, s, meta, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"], size=26)
    has_bul = bool(s.get("bullets"))
    cap_h = Inches(0.55) if s.get("caption") else 0
    fig_w = Inches(7.6) if has_bul else W - 2 * MARGIN
    fig_h = CONTENT_BOTTOM - CONTENT_TOP - cap_h - Inches(0.3)
    place_image(sl, fig_path(s["figure"]), MARGIN, CONTENT_TOP + Inches(0.05), fig_w, fig_h)
    if s.get("caption"):
        add_text(sl, MARGIN, CONTENT_BOTTOM - cap_h - Inches(0.2), fig_w, cap_h, s["caption"], size=13, color=GREY, italic=True)
    if has_bul:
        left = MARGIN + fig_w + Inches(0.4); bw = W - MARGIN - left
        size = fit_size(s["bullets"], 18, bw / 914400, fig_h / 914400, line_h=1.6)
        add_bullets(sl, left, CONTENT_TOP + Inches(0.2), bw, fig_h, s["bullets"], size=size, space_after=10)
    credit(sl, meta); footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_figures(prs, s, meta, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"], size=26)
    cap_h = Inches(0.55) if s.get("caption") else 0
    avail = CONTENT_BOTTOM - CONTENT_TOP - cap_h - Inches(0.3); k = len(s["figures"]); gap = Inches(0.1)
    each = (avail - gap * (k - 1)) // k; top = CONTENT_TOP + Inches(0.05)
    for slug in s["figures"]:
        place_image(sl, fig_path(slug), MARGIN, top, W - 2 * MARGIN, each); top += each + gap
    if s.get("caption"):
        add_text(sl, MARGIN, CONTENT_BOTTOM - cap_h - Inches(0.2), W - 2 * MARGIN, cap_h, s["caption"], size=13, color=GREY, italic=True)
    credit(sl, meta); footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_table(prs, s, meta, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"], size=26)
    cols, rows = s["columns"], s["rows"]; nr, nc = len(rows) + 1, len(cols)
    box_h = CONTENT_BOTTOM - CONTENT_TOP - Inches(0.3)
    shape = sl.shapes.add_table(nr, nc, MARGIN, CONTENT_TOP + Inches(0.15), W - 2 * MARGIN, box_h)
    tbl = shape.table
    # column widths: first column narrower when it is a label column
    total = W - 2 * MARGIN
    if nc == 3 and cols[0] == "": widths = [0.22, 0.39, 0.39]
    elif nc == 3: widths = [0.42, 0.29, 0.29]
    elif nc == 4: widths = [0.18, 0.26, 0.28, 0.28]
    else: widths = [1 / nc] * nc
    for i, wfrac in enumerate(widths): tbl.columns[i].width = int(total * wfrac)
    longest = max(len(c) for r in rows for c in r)
    size = 15 if nr <= 6 else 13 if nr <= 8 else 12
    if longest > 70: size -= 1
    for c, name in enumerate(cols):
        cell = tbl.cell(0, c); cell.text = name; cell.fill.solid(); cell.fill.fore_color.rgb = LIGHT
        for p in cell.text_frame.paragraphs:
            for r in p.runs: r.font.bold = True; r.font.size = Pt(size); r.font.name = FONT; rgb(r.font, BLACK)
    for ri, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.cell(ri, c); cell.text = val; cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size); r.font.name = FONT; rgb(r.font, BLACK); r.font.bold = (c == 0 and cols[0] == "")
    footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_two_col(prs, s, meta, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"])
    corner_logo(sl, s)
    col_w = (W - 2 * MARGIN - Inches(0.5)) // 2; box_h = CONTENT_BOTTOM - CONTENT_TOP - Inches(0.7)
    for i, key in enumerate(("left", "right")):
        col = s[key]; left = MARGIN + i * (col_w + Inches(0.5))
        add_text(sl, left, CONTENT_TOP + Inches(0.1), col_w, Inches(0.5), col["heading"], size=20, bold=True)
        rule(sl, CONTENT_TOP + Inches(0.62), left=left, width=col_w, color=ACCENT if i == 0 else GREY, h=Pt(2))
        size = fit_size(col["bullets"], 18, col_w / 914400, box_h / 914400, line_h=1.6)
        add_bullets(sl, left, CONTENT_TOP + Inches(0.75), col_w, box_h, col["bullets"], size=size, space_after=10)
    footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_callout(prs, s, meta, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"])
    top = CONTENT_TOP + Inches(0.6); h = Inches(3.2)
    bar = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, top, Pt(6), h)
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background(); bar.shadow.inherit = False
    size = fit_size([s["text"]], 28, (W - 2 * MARGIN - Inches(0.5)) / 914400, h / 914400, line_h=1.4)
    add_text(sl, MARGIN + Inches(0.4), top, W - 2 * MARGIN - Inches(0.4), h, s["text"], size=size, anchor=MSO_ANCHOR.MIDDLE)
    footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_prompt(prs, s, meta, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"])
    lines = s["text"].rstrip("\n").split("\n"); box_h = CONTENT_BOTTOM - CONTENT_TOP - Inches(0.3)
    box = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, CONTENT_TOP + Inches(0.15), W - 2 * MARGIN, box_h)
    box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.color.rgb = RULE; box.shadow.inherit = False
    size = fit_size(lines, 16, (W - 2 * MARGIN - Inches(0.6)) / 914400 * 0.85, (box_h - Inches(0.4)) / 914400, line_h=1.3)
    tb = add_text(sl, MARGIN + Inches(0.3), CONTENT_TOP + Inches(0.35), W - 2 * MARGIN - Inches(0.6), box_h - Inches(0.4), lines,
                  size=size, font=MONO)
    footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_steps(prs, s, meta, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"])
    steps = s["steps"]; k = len(steps); avail = CONTENT_BOTTOM - CONTENT_TOP - Inches(0.2); each = avail // k
    hsize = 18 if k <= 5 else 16; tsize = 15 if k <= 5 else 13
    for i, st in enumerate(steps):
        top = CONTENT_TOP + Inches(0.1) + each * i
        dot = sl.shapes.add_shape(MSO_SHAPE.OVAL, MARGIN, top + Inches(0.08), Inches(0.42), Inches(0.42))
        dot.fill.solid(); dot.fill.fore_color.rgb = ACCENT; dot.line.fill.background(); dot.shadow.inherit = False
        dot.text_frame.text = str(i + 1); p = dot.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size = Pt(13); p.runs[0].font.bold = True; rgb(p.runs[0].font, WHITE); p.runs[0].font.name = FONT
        dot.text_frame.margin_left = dot.text_frame.margin_right = 0
        add_text(sl, MARGIN + Inches(0.6), top, W - 2 * MARGIN - Inches(0.6), Inches(0.5), st["heading"], size=hsize, bold=True)
        add_text(sl, MARGIN + Inches(0.6), top + Inches(0.45), W - 2 * MARGIN - Inches(0.6), each - Inches(0.5), st["text"], size=tsize)
        if i < k - 1: rule(sl, top + each - Inches(0.05), left=MARGIN + Inches(0.6), width=W - 2 * MARGIN - Inches(0.6))
    footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_quiz(prs, s, meta, n, reveal=False):
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"] + (" · answer" if reveal else ""))
    add_text(sl, MARGIN, CONTENT_TOP + Inches(0.1), W - 2 * MARGIN, Inches(1.2), s["question"], size=22, bold=True)
    top = CONTENT_TOP + Inches(1.4); letters = "ABCD"
    for i, opt in enumerate(s["options"]):
        is_ans = reveal and i == s["answer"]
        box = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, top, W - 2 * MARGIN, Inches(0.62))
        box.fill.solid(); box.fill.fore_color.rgb = LIGHT if not is_ans else WHITE
        box.line.color.rgb = ACCENT if is_ans else RULE; box.line.width = Pt(2.5 if is_ans else 1); box.shadow.inherit = False
        add_text(sl, MARGIN + Inches(0.2), top + Inches(0.08), W - 2 * MARGIN - Inches(0.4), Inches(0.5),
                 f"{letters[i]}.  {opt}", size=17, bold=is_ans, anchor=MSO_ANCHOR.MIDDLE)
        top += Inches(0.72)
    if reveal:
        add_text(sl, MARGIN, top + Inches(0.1), W - 2 * MARGIN, Inches(0.9), "Why: " + s["why"], size=15, color=GREY, italic=True)
    footer(sl, meta, n)
    notes(sl, (s.get("notes") or "") + f"\nCorrect answer: {letters[s['answer']]}. {s['why']}")
    return sl


def l_references(prs, s, meta, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"])
    refs = s.get("refs") or meta["references"]
    size = fit_size(refs, 14, (W - 2 * MARGIN) / 914400, (CONTENT_BOTTOM - CONTENT_TOP - Inches(0.3)) / 914400, line_h=1.5)
    add_bullets(sl, MARGIN, CONTENT_TOP + Inches(0.15), W - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP - Inches(0.3), refs,
                size=size, space_after=6, numbered=True)
    footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_articles(prs, s, meta, n):
    """Latest-articles slide: each item has title, source, date, url, takeaway."""
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"])
    items = s["items"]; k = len(items); avail = CONTENT_BOTTOM - CONTENT_TOP - Inches(0.2); each = avail // k
    for i, it in enumerate(items):
        top = CONTENT_TOP + Inches(0.1) + each * i
        add_text(sl, MARGIN, top, W - 2 * MARGIN, Inches(0.45), it["title"], size=16 if k <= 4 else 14, bold=True)
        add_text(sl, MARGIN, top + Inches(0.4), W - 2 * MARGIN, Inches(0.35), f"{it['source']} · {it['date']} · {it['url']}",
                 size=11, color=GREY)
        add_text(sl, MARGIN, top + Inches(0.7), W - 2 * MARGIN, each - Inches(0.75), it["takeaway"], size=13 if k <= 4 else 12)
        if i < k - 1: rule(sl, top + each - Inches(0.05))
    footer(sl, meta, n); notes(sl, s.get("notes")); return sl


def l_qr(prs, s, meta, n):
    """Join slide: QR image (s['qr_image'] path, optional) or a placeholder box, bot handle, rules."""
    sl = prs.slides.add_slide(prs.slide_layouts[6]); title_bar(sl, s["title"])
    box = Inches(3.6); left = MARGIN; top = CONTENT_TOP + Inches(0.3)
    img = s.get("qr_image")
    if img and (ROOT / img).exists():
        place_image(sl, ROOT / img, left, top, box, box)
    else:
        ph = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, box, box)
        ph.fill.solid(); ph.fill.fore_color.rgb = LIGHT; ph.line.color.rgb = GREY; ph.line.dash_style = 4; ph.shadow.inherit = False
        add_text(sl, left, top + box // 2 - Inches(0.5), box, Inches(1.0), s.get("qr_placeholder", "Paste the bot's QR code here"),
                 size=16, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(sl, left, top + box + Inches(0.15), box, Inches(0.5), s.get("handle", ""), size=18, bold=True, align=PP_ALIGN.CENTER)
    rleft = left + box + Inches(0.6); rw = W - MARGIN - rleft
    add_text(sl, rleft, top, rw, Inches(0.5), s.get("steps_heading", "How to join"), size=20, bold=True)
    add_bullets(sl, rleft, top + Inches(0.55), rw, Inches(2.0), s["steps"], size=17, numbered=True, space_after=8)
    add_text(sl, rleft, top + Inches(2.5), rw, Inches(0.5), s.get("rules_heading", "Three rules"), size=20, bold=True)
    add_bullets(sl, rleft, top + Inches(3.05), rw, Inches(1.8), s["rules"], size=16, space_after=6)
    footer(sl, meta, n); notes(sl, s.get("notes")); return sl


LAYOUTS = {"title": l_title, "section": l_section, "bullets": l_bullets, "figure": l_figure, "figures": l_figures,
           "table": l_table, "two_col": l_two_col, "callout": l_callout, "prompt": l_prompt, "steps": l_steps,
           "references": l_references, "articles": l_articles, "qr": l_qr}


def build(yaml_path, out_path):
    doc = yaml.safe_load(pathlib.Path(yaml_path).read_text()); meta = doc["meta"]
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    n = 0
    for s in doc["slides"]:
        n += 1
        if s["layout"] == "quiz":
            l_quiz(prs, s, meta, n); n += 1; l_quiz(prs, s, meta, n, reveal=True)
        else:
            LAYOUTS[s["layout"]](prs, s, meta, n)
    pathlib.Path(out_path).parent.mkdir(parents=True, exist_ok=True); prs.save(out_path)
    print(f"{out_path}: {n} slides, {len(used_figures)} figures")
    return used_figures


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
